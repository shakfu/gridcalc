"""Write an xlsx of formulas whose Excel result is unconfirmed, with gridcalc's answers.

A:H hold the inputs every case reads; each case row holds the formula as text,
the live formula, the gridcalc result and a match flag. Open it in Excel 365.
Procedure: `docs/dev/excel-check.md`.

    uv run python scripts/excel_check.py build/excel-check.xlsx
"""

from __future__ import annotations

import argparse
import math
import re
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula

from gridcalc.cli import run_eval
from gridcalc.engine import Grid, Mode, _add_xl_prefixes

# (cell, entry). A leading `'` marks text, as typing it in Excel does.
INPUTS = [
    ("A1", "5"), ("A3", "7"), ("A4", "-2.5"),  # A2 blank
    ("B1", "1"), ("B2", "2"), ("B3", "'3"), ("B4", "'a"),
    ("C1", "10"), ("C2", "20"), ("C3", "30"), ("C4", "40"),
    ("D1", "'a"), ("D2", "'b"), ("D3", "'a"), ("D4", "'c"),
    ("F1", "'Name"), ("G1", "'Val"),
    ("F2", "'x"), ("G2", "1"),
    ("F3", "'y"), ("G3", "=NA()"),
    ("F4", "'x"), ("G4", "3"),
    ("H1", "'Name"), ("H2", "'x"),
    # A second database whose matched row holds the error.
    ("F6", "'Name"), ("G6", "'Val"),
    ("F7", "'x"), ("G7", "=NA()"),
    ("F8", "'y"), ("G8", "5"),
    ("H6", "'Name"), ("H7", "'x"),
]  # fmt: skip

# (id, formula, what is in doubt). Array results are joined to one cell, so
# nothing spills into the next case.
CASES = [
    ("setup", "=AND(ISTEXT(B3),ISBLANK(A2),ISNA(G3))", "TRUE means the inputs pasted as intended"),
    ("pow-1", "=2^3^2", "Excel documents left to right (64)"),
    ("pow-2", "=-2^2", "negation binds tighter than ^ in Excel (4)"),
    ("pow-3", "=-2^0.5", "(-2)^0.5 in Excel: #NUM!"),
    ("text-1", '=0.00001&""', "gridcalc formats with %g"),
    ("text-2", '=1E-10&""', ""),
    ("text-3", '=-0.0000123&""', ""),
    ("text-4", '=1/3&""', "15 significant digits"),
    ("text-5", '=(0.1+0.2)&""', ""),
    ("text-6", '=123456789012&""', ""),
    ("text-7", '=12345678901234567&""', "past 15 digits"),
    ("text-8", '=1E+20&""', ""),
    ("text-9", '=1E-11&""', "gridcalc writes exponents below -10 in scientific form; edge unknown"),
    ("text-10", '=1E-15&""', ""),
    ("text-11", '=1E-19&""', ""),
    ("text-12", '=1/3*1E-10&""', "15 digits past the edge"),
    ("text-13", '=1E+19&""', "gridcalc writes exponents 20 and up in scientific form"),
    ("text-14", '=2^53&""', "16-digit integer: rounded to 15 digits, or scientific?"),
    ("text-15", '=1E-18&""', "plain text of 20 characters: plain if the edge is 20 characters"),
    ("text-16", '=-1E-18&""', "21 characters with the sign"),
    ("text-17", '=1.5E-17&""', "20 characters"),
    ("text-18", '=1.5E+19&""', "20 characters"),
    ("text-19", '=1.2345E-15&""', "21 characters, exponent -15"),
    ("num-1", "=12345678901234567-12345678901234500", "Excel keeps 15 digits of a typed number"),
    ("blank-1", "=ROUND(2.5,)", "omitted number argument reads as 0"),
    ("blank-2", "=ROUND(2.5,Z9)", "blank cell reads as 0"),
    ("blank-3", '=LEFT("abc",)', "omitted count reads as 0"),
    ("blank-4", '=SUBSTITUTE("abc","b",)', "omitted text reads as empty"),
    ("blank-5", '=CONCAT("a",,"b")', ""),
    ("blank-6", "=IF(FALSE,1,)", ""),
    ("lift-1", '=TEXTJOIN(",",FALSE,LEN(A1:A3))', "scalar function over a range"),
    ("lift-2", '=TEXTJOIN(",",FALSE,ROUND(A1:A4,0))', "ROUND(-2.5,0) rounds away from zero"),
    ("lift-3", '=TEXTJOIN(",",FALSE,MOD(C1:C4,3))', ""),
    ("lift-4", '=TEXTJOIN(",",FALSE,UPPER(D1:D4))', ""),
    ("lift-5", '=TEXTJOIN(",",FALSE,TEXT(C1:C3,"0.0"))', ""),
    ("lift-6", "=SUM(LEN(A1:A3))", ""),
    ("crit-1", "=COUNTIF(B1:B3,3)", "does the text '3 count?"),
    ("crit-2", '=COUNTIF(B1:B3,"3")', ""),
    ("crit-3", '=COUNTIF(B1:B3,">1")', ""),
    ("crit-4", '=SUMIF(B1:B4,">0")', ""),
    ("crit-5", '=SUMIF(C1:C4,">15",A1:A2)', "Excel resizes sum_range from its top-left"),
    ("crit-6", '=TEXTJOIN(",",FALSE,SUMIF(D1:D4,{"a","b"},C1:C4))', "array criteria"),
    ("crit-7", '=SUM(COUNTIF(D1:D4,{"a","c"}))', ""),
    ("crit-8", '=COUNTIF(B1:B4,"<>3")', "gridcalc counts only numbers for <>"),
    ("crit-9", '=COUNTIF(A1:A4,"<>5")', "does the blank A2 count?"),
    ("arr-1", "=ROWS(FILTER(C1:C4,C1:C4>15))", ""),
    ("arr-2", '=TEXTJOIN(",",FALSE,FREQUENCY({1,2,3,4,5},{4,2}))', "unsorted bins"),
    ("arr-3", '=TEXTJOIN(",",FALSE,ISFORMULA(G1:G3))', "ISFORMULA over a range"),
    ("db-1", '=DSUM(F1:G4,"Val",H1:H2)', "an error in an unmatched row"),
    ("db-2", '=DCOUNT(F1:G4,"Val",H1:H2)', ""),
    ("db-3", '=DSUM(F6:G8,"Val",H6:H7)', "an error in the matched row: gridcalc propagates it"),
    ("db-4", '=DCOUNT(F6:G8,"Val",H6:H7)', ""),
    ("date-1", "=EOMONTH(DATE(1900,1,15),1)", "phantom 1900-02-29"),
    ("date-2", "=DAYS(10.9,1.1)", "fractional serials"),
    ("yf-1", "=YEARFRAC(DATE(2023,12,31),DATE(2024,3,1),0)", ""),
    ("yf-2", "=YEARFRAC(DATE(2023,12,31),DATE(2024,3,1),1)", ""),
    ("yf-3", "=YEARFRAC(DATE(2023,6,15),DATE(2025,2,28),0)", ""),
    ("yf-4", "=YEARFRAC(DATE(2023,6,15),DATE(2025,2,28),1)", ""),
    ("yf-5", "=YEARFRAC(DATE(2024,2,29),DATE(2025,2,28),0)", ""),
    ("yf-6", "=YEARFRAC(DATE(2024,2,29),DATE(2025,2,28),1)", ""),
    ("yf-7", "=YEARFRAC(DATE(2023,2,28),DATE(2024,2,29),0)", ""),
    ("yf-8", "=YEARFRAC(DATE(2023,2,28),DATE(2024,2,29),1)", ""),
    # Basis 1 under a year apart: is the year 366 days when the span holds
    # Feb 29, when either date's year is a leap year, or never?
    ("yf-9", "=YEARFRAC(DATE(2024,1,1),DATE(2024,1,31),1)", "leap year, no Feb 29 in span"),
    ("yf-10", "=YEARFRAC(DATE(2023,3,1),DATE(2024,2,28),1)", "ends in a leap year before Feb 29"),
    ("yf-11", "=YEARFRAC(DATE(2023,3,1),DATE(2024,2,29),1)", "ends on Feb 29"),
    ("yf-12", "=YEARFRAC(DATE(2024,3,1),DATE(2025,2,28),1)", "starts in a leap year after Feb 29"),
    ("yf-13", "=YEARFRAC(DATE(2023,1,1),DATE(2023,12,31),1)", "no leap year"),
    ("yf-14", "=YEARFRAC(DATE(2024,1,1),DATE(2025,1,1),1)", "exactly one leap year"),
    ("yf-15", "=YEARFRAC(DATE(2023,1,1),DATE(2024,1,2),1)", "just over a year: averaged"),
    # Basis 0, US 30/360 day rules.
    ("yf-16", "=YEARFRAC(DATE(2024,1,31),DATE(2024,3,31),0)", "both on the 31st"),
    ("yf-17", "=YEARFRAC(DATE(2024,2,28),DATE(2024,3,31),0)", "leap-year Feb 28: not month end"),
    ("yf-18", "=YEARFRAC(DATE(2023,2,28),DATE(2023,3,31),0)", "starts on the last day of Feb"),
    ("yf-19", "=YEARFRAC(DATE(2024,2,29),DATE(2024,3,30),0)", ""),
    ("yf-20", "=YEARFRAC(DATE(2023,1,30),DATE(2023,2,28),0)", "ends on the last day of Feb"),
    ("misc-1", "=MOD(3,0.1)", "float residue"),
    ("misc-2", '=IMSQRT("-4")', ""),
    ("misc-3", '=OR("a",TRUE)', "text literal in OR"),
    ("misc-4", '=XOR(TRUE,"a")', ""),
    ("misc-5", "=MATCH(3,{5,1,4,2},1)", "approximate match over unsorted data"),
    ("misc-6", "=TEXT(3,1)", ""),
    ("misc-7", '=OR("a",TRUE)', "misc-3 as a plain formula"),
    ("misc-8", '=XOR(TRUE,"a")', "misc-4 as a plain formula"),
    ("logic-1", '=AND("a",TRUE)', "gridcalc: #VALUE!; OR skips the text"),
    ("logic-2", '=OR("a")', "no logical value left: #VALUE!?"),
    ("logic-3", '=OR("TRUE",FALSE)', "text TRUE read as a logical?"),
    ("3d-1", "=SUM(S3a:S3c!B2)", "hand check after the run: see excel-check.md"),
    ("3d-2", "=COUNTA(S3a:S3c!B2:B3)", "text on one spanned sheet counts"),
    (
        "accrint-1",
        "=ACCRINT(DATE(2008,3,1),DATE(2008,8,31),DATE(2008,5,1),0.1,1000,2,1)",
        "basis 1: quasi-coupon periods anchored at first_interest (inference)",
    ),
    (
        "accrint-2",
        "=ACCRINT(DATE(2008,3,1),DATE(2008,8,31),DATE(2008,9,30),0.1,1000,2,1)",
        "basis 1 across a coupon date",
    ),
    (
        "chisq-1",
        "=CHISQ.TEST({58,35;11,25;10,23},{45.35,47.65;17.56,18.44;16.09,16.91})",
        "2D table: df = (rows-1)(cols-1)",
    ),
    ("chisq-2", "=CHISQ.TEST({1,2,3},{1,2})", "different sizes: #N/A per Microsoft"),
    ("ind-1", '=INDIRECT("A"&3)+INDIRECT("R3C1",FALSE)', ""),
    ("whole-1", "=COUNTBLANK(Z:Z)", "gridcalc reads Z:Z to the used extent"),
    ("whole-2", '=COUNTIF(Z:Z,"<>3")', "blanks past the used extent count"),
]

# Sheets after `check` for the 3d-* cases: B2 on each, and text in S3b!B3.
SPAN_SHEETS = [("S3a", {"B2": "10"}), ("S3b", {"B2": "20", "B3": "'x"}), ("S3c", {"B2": "30"})]

# Cases stored as plain formulas: twins of array-formula cases, and 3D
# references, which Excel refuses in an array formula.
PLAIN = frozenset({"misc-7", "misc-8", "3d-1", "3d-2"})

# I id, J formula as text, K live formula, L gridcalc, M match, N note.
# `=` ignores case on text, so text goes through EXACT.
MATCH = (
    "=IFERROR(IF(AND(ISTEXT(K{r}),ISTEXT(L{r})),EXACT(K{r},L{r}),K{r}=L{r}),"
    "IFERROR(ERROR.TYPE(K{r})=ERROR.TYPE(L{r}),FALSE))"
)


def _cell(ref: str) -> tuple[int, int]:
    return ord(ref[0]) - ord("A"), int(ref[1:]) - 1


def _result(entry: dict[str, object]) -> object:
    """gridcalc's answer as an xlsx value; openpyxl stores an error code as an error."""
    if entry["error"]:
        return str(entry["error"])
    v = entry["value"]
    if v in ("TRUE", "FALSE"):  # how run_eval reports a logical; no case returns that text
        return v == "TRUE"
    if isinstance(v, float) and math.isnan(v):
        return None
    return v


def evaluate() -> list[object]:
    """gridcalc's result for each case, in ``CASES`` order."""
    g = Grid()
    g.mode = Mode.EXCEL
    g._apply_mode_libs()
    for ref, entry in INPUTS:
        c, r = _cell(ref)
        if entry.startswith("'"):
            g._put_label(c, r, entry[1:])
        else:
            g.setcell(c, r, entry)
    for i, (name, cells) in enumerate(SPAN_SHEETS, start=1):
        g.add_sheet(name)
        g.set_active(i)
        for ref, entry in cells.items():
            c, r = _cell(ref)
            if entry.startswith("'"):
                g._put_label(c, r, entry[1:])
            else:
                g.setcell(c, r, entry)
    g.set_active(0)
    g.recalc()
    return [_result(e) for e in run_eval(g, [f for _, f, _ in CASES])]


def build() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "check"
    for ref, entry in INPUTS:
        ws[ref] = entry[1:] if entry.startswith("'") else _number_or_formula(entry)
    heads = ["id", "formula", "Excel", "gridcalc", "match", "in doubt"]
    for col, head in zip("IJKLMN", heads, strict=True):
        ws[f"{col}1"] = head
    for r, ((cid, formula, note), res) in enumerate(zip(CASES, evaluate(), strict=True), start=2):
        ws[f"I{r}"] = cid
        ws[f"J{r}"] = formula
        ws[f"J{r}"].data_type = "s"  # shown as text, not computed
        # An array formula, so a range reaching a one-value parameter is not
        # reduced to one cell by implicit intersection.
        stored = _add_xl_prefixes(formula)
        ws[f"K{r}"] = stored if cid in PLAIN else ArrayFormula(f"K{r}", stored)
        ws[f"L{r}"] = res
        ws[f"M{r}"] = MATCH.format(r=r)
        ws[f"N{r}"] = note
    for name, cells in SPAN_SHEETS:
        sheet = wb.create_sheet(name)
        for ref, entry in cells.items():
            sheet[ref] = entry[1:] if entry.startswith("'") else float(entry)
    wb.calculation.fullCalcOnLoad = True  # openpyxl writes no cached values
    return wb


def _number_or_formula(entry: str) -> object:
    return entry if entry.startswith("=") else float(entry)


def save(wb: Workbook, path: Path) -> None:
    """Save ``wb`` with each number in L stored to 17 digits.

    openpyxl writes 16 significant digits, which can lose the last bit. A
    formula literal is no fix: Excel keeps 15 digits of a typed number.
    """
    wb.save(path)
    ws = wb.active
    exact = {
        f"L{r}": repr(ws[f"L{r}"].value)
        for r in range(2, ws.max_row + 1)
        if isinstance(ws[f"L{r}"].value, float)
    }
    part = "xl/worksheets/sheet1.xml"
    with zipfile.ZipFile(path) as z:
        members = [(info, z.read(info)) for info in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in members:
            if info.filename == part:
                data = re.sub(
                    rb'(<c r="(L\d+)" t="n"><v>)[^<]*(</v>)',
                    lambda m: m[1] + exact[m[2].decode()].encode() + m[3],
                    data,
                )
            z.writestr(info, data)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("out", type=Path)
    args = parser.parse_args()
    save(build(), args.out)


if __name__ == "__main__":
    main()
