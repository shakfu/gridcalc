"""Compare gridcalc's results with the values Excel saved in real workbooks.

Excel stores each formula's last result next to the formula. This script reads
those values, recalculates the workbook in gridcalc, and reports the formula
cells that disagree. A mismatch whose inputs all agree is a root; one that
reads a mismatched cell is downstream, and is counted but not listed.

    uv run python scripts/excel_corpus.py book1.xlsx book2.xlsx
    uv run python scripts/excel_corpus.py --json out.json corpus/*.xlsx
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import re
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook as openpyxl_load
from openpyxl.utils.datetime import to_excel

from gridcalc.engine import FORMULA, Grid
from gridcalc.formula.deps import extract_refs
from gridcalc.loader import load_workbook

Key = tuple[str, int, int]  # (sheet, col, row), 0-based
_FUNC_RE = re.compile(r"\b([A-Z][A-Z0-9.]*)\(")


def application(path: Path) -> str:
    """The `Application` a file's `docProps/app.xml` names, or ""."""
    try:
        with zipfile.ZipFile(path) as z:
            root = ET.fromstring(z.read("docProps/app.xml"))  # noqa: S314
    except (KeyError, OSError, zipfile.BadZipFile, ET.ParseError):
        return ""
    for el in root:
        if el.tag.endswith("}Application"):
            return el.text or ""
    return ""


def excel_values(path: Path) -> tuple[dict[Key, str], dict[Key, Any]]:
    """Formula text and Excel's saved value for each formula cell."""
    formulas: dict[Key, str] = {}
    for ws in openpyxl_load(path, data_only=False):
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type != "f":  # text that starts with "=" is not a formula
                    continue
                text = getattr(cell.value, "text", cell.value)  # ArrayFormula holds its text
                formulas[(ws.title, cell.column - 1, cell.row - 1)] = str(text)
    saved: dict[Key, Any] = {}
    for ws in openpyxl_load(path, data_only=True):
        for row in ws.iter_rows():
            for cell in row:
                key = (ws.title, cell.column - 1, cell.row - 1)
                if key not in formulas or cell.value is None:
                    continue
                v = cell.value
                if cell.data_type == "e":
                    v = ("error", str(v))
                elif isinstance(v, (dt.datetime, dt.date, dt.time)):
                    v = float(to_excel(v))
                saved[key] = v
    return formulas, saved


def gridcalc_value(g: Grid, key: Key) -> Any:
    """gridcalc's result for a cell, in the form `excel_values` uses."""
    sheet = next(s for s in g.sheets if s.name == key[0])
    cl = sheet._cells.get((key[1], key[2]))
    if cl is None:
        return None
    if cl.err is not None:
        return ("error", str(cl.err))
    if isinstance(cl.val, bool):
        return cl.val
    if cl.sval is not None:
        return cl.sval
    if isinstance(cl.val, float) and math.isnan(cl.val):
        return ("error", "ERROR")
    return cl.val


def same(a: Any, b: Any, rel: float) -> bool:
    if isinstance(a, bool) or isinstance(b, bool):
        return a is b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return math.isclose(a, b, rel_tol=rel, abs_tol=1e-12)
    return bool(a == b)


def functions(formula: str) -> set[str]:
    return {
        m.removeprefix("_XLFN.").removeprefix("_XLWS.") for m in _FUNC_RE.findall(formula.upper())
    }


def compare(path: Path, rel: float = 1e-9) -> dict[str, Any]:
    """Compare one workbook. Returns counts, root mismatches and their functions."""
    report: dict[str, Any] = {"file": str(path), "application": application(path)}
    try:
        formulas, saved = excel_values(path)
        g = load_workbook(path)
    except Exception as exc:  # noqa: BLE001 -- one bad file must not stop the run
        report["error"] = f"{type(exc).__name__}: {exc}"
        return report
    as_value: list[Key] = []
    bad: dict[Key, tuple[Any, Any]] = {}
    for key in formulas:
        if key not in saved:
            continue
        sheet = next((s for s in g.sheets if s.name == key[0]), None)
        cl = sheet._cells.get((key[1], key[2])) if sheet is not None else None
        if cl is None or cl.type != FORMULA:
            as_value.append(key)
            continue
        got = gridcalc_value(g, key)
        if not same(saved[key], got, rel):
            bad[key] = (saved[key], got)
    roots = []
    for key in bad:
        cl = next(s for s in g.sheets if s.name == key[0])._cells[(key[1], key[2])]
        reads = extract_refs(cl.ast, formula_sheet=key[0]) if cl.ast is not None else set()
        if not any((s or key[0], c, r) in bad for s, c, r in reads):
            roots.append(key)
    by_function: Counter[str] = Counter()
    for key in roots:
        by_function.update(functions(formulas[key]))
    compared = len(saved) - len(as_value)
    report.update(
        formulas=len(formulas),
        no_saved_value=len(formulas) - len(saved),
        imported_as_value=len(as_value),
        compared=compared,
        agree=compared - len(bad),
        roots=[
            {"cell": f"{s}!{_a1(c, r)}", "formula": formulas[(s, c, r)], "excel": e, "gridcalc": v}
            for (s, c, r) in roots
            for e, v in [bad[(s, c, r)]]
        ],
        downstream=len(bad) - len(roots),
        by_function=dict(by_function.most_common()),
    )
    return report


def _a1(c: int, r: int) -> str:
    letters = ""
    c += 1
    while c:
        c, rem = divmod(c - 1, 26)
        letters = chr(65 + rem) + letters
    return f"{letters}{r + 1}"


def _show(v: Any) -> str:
    return v[1] if isinstance(v, tuple) else repr(v)


def render(reports: list[dict[str, Any]], limit: int) -> str:
    lines: list[str] = []
    total: Counter[str] = Counter()
    for rep in reports:
        lines.append(f"== {rep['file']}  ({rep['application'] or 'unknown application'})")
        if "error" in rep:
            lines.append(f"   not read: {rep['error']}")
            continue
        if rep["formulas"] and rep["no_saved_value"] == rep["formulas"]:
            lines.append("   no saved values: the file was not last saved by Excel")
            continue
        lines.append(
            f"   {rep['agree']} of {rep['compared']} agree; {len(rep['roots'])} root and "
            f"{rep['downstream']} downstream mismatches; {rep['imported_as_value']} imported "
            f"as values; {rep['no_saved_value']} without a saved value"
        )
        for m in rep["roots"][:limit]:
            lines.append(
                f"   {m['cell']}: {m['formula']}  excel={_show(m['excel'])}  "
                f"gridcalc={_show(m['gridcalc'])}"
            )
        if len(rep["roots"]) > limit:
            lines.append(f"   ... {len(rep['roots']) - limit} more")
        total.update(rep["by_function"])
    if total:
        lines.append("== root mismatches by function")
        lines += [f"   {n:5d}  {name}" for name, n in total.most_common()]
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("files", nargs="+", type=Path)
    parser.add_argument("--json", type=Path, help="write the full reports here")
    parser.add_argument("--rel", type=float, default=1e-9, help="relative number tolerance")
    parser.add_argument("--limit", type=int, default=20, help="root mismatches listed per file")
    args = parser.parse_args()
    reports = [compare(p, args.rel) for p in args.files]
    if args.json:
        args.json.write_text(json.dumps(reports, indent=2, default=str))
    print(render(reports, args.limit))


if __name__ == "__main__":
    main()
