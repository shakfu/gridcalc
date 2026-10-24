"""Cell styles beyond bold/italic/underline: fonts, fills, borders, alignment."""

import json
from pathlib import Path

import pytest

from gridcalc.cellstyle import Side, XStyle
from gridcalc.engine import EMPTY, NCOL, NROW, SPILL, Grid, Mode

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill  # noqa: E402
from openpyxl.styles import Side as XSide  # noqa: E402

POSITIONS = [(0, 0), (1, 0), (2, 2), (3, 3), (4, 4)]


def _excel() -> Grid:
    g = Grid()
    g.mode = Mode.EXCEL
    g._apply_mode_libs()
    return g


def _rich(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "head"
    ws["A1"].font = Font(name="Arial", size=14, color="FF0000", bold=True, strike=True)
    ws["B1"] = 5
    ws["B1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["C3"].border = Border(top=XSide("thin", color="0000FF"), bottom=XSide("double"))
    ws["D4"] = 2
    ws["D4"].alignment = Alignment(horizontal="center", vertical="top", wrap_text=True, indent=1)
    ws["E5"].fill = PatternFill("solid", fgColor="00FF00")
    wb.save(str(path))


def test_xlsx_import_reads_fonts_fills_borders_and_alignment(tmp_path) -> None:
    _rich(tmp_path / "rich.xlsx")
    g = Grid()
    assert g.xlsxload(str(tmp_path / "rich.xlsx")) == 0
    assert g.cell(0, 0).bold == 1
    assert g.cell(0, 0).xstyle == XStyle(font="Arial", size=14.0, color="FF0000", strike=True)
    assert g.cell(1, 0).xstyle == XStyle(fill="FFFF00")
    assert g.cell(2, 2).type == EMPTY  # a styled blank is kept
    assert g.cell(2, 2).xstyle == XStyle(top=Side("thin", "0000FF"), bottom=Side("double"))
    assert g.cell(3, 3).xstyle == XStyle(halign="center", valign="top", wrap=True, indent=1)


def test_xlsx_round_trip(tmp_path) -> None:
    _rich(tmp_path / "rich.xlsx")
    g = Grid()
    g.xlsxload(str(tmp_path / "rich.xlsx"))
    g.mode = Mode.EXCEL
    assert g.xlsxsave(str(tmp_path / "out.xlsx")) == 0
    ws = openpyxl.load_workbook(str(tmp_path / "out.xlsx")).active
    assert (ws["A1"].font.name, ws["A1"].font.sz, ws["A1"].font.b) == ("Arial", 14, True)
    assert ws["B1"].fill.fgColor.rgb == "FFFFFF00"
    assert (ws["C3"].border.top.style, ws["C3"].border.bottom.style) == ("thin", "double")
    assert ws["D4"].alignment.horizontal == "center" and ws["D4"].alignment.wrap_text
    assert ws["E5"].fill.fgColor.rgb == "FF00FF00"
    g2 = Grid()
    g2.xlsxload(str(tmp_path / "out.xlsx"))
    assert [g2.cell(*p).xstyle for p in POSITIONS] == [g.cell(*p).xstyle for p in POSITIONS]


def test_theme_colours_resolve_with_tint(tmp_path) -> None:
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    # openpyxl's default theme is Office 2007's: accent1 (index 4) is 4F81BD.
    wb.active["A1"].fill = PatternFill("solid", fgColor=Color(theme=4))
    wb.active["A2"] = 1
    wb.active["A2"].fill = PatternFill("solid", fgColor=Color(theme=4, tint=0.4))
    wb.save(str(tmp_path / "t.xlsx"))
    g = Grid()
    g.xlsxload(str(tmp_path / "t.xlsx"))
    assert g.cell(0, 0).xstyle.fill == "4F81BD"
    assert g.cell(0, 1).xstyle.fill == "95B3D7"  # Excel's "Lighter 40%" swatch


def test_json_round_trip_keeps_styles_and_styled_blanks(tmp_path) -> None:
    g = _excel()
    g.setcell(0, 0, "5")
    g.cell(0, 0).xstyle = XStyle(fill="FFFF00", left=Side("thick", "FF0000"))
    g._ensure_cell(1, 1).xstyle = XStyle(halign="center")
    f = tmp_path / "b.json"
    assert g.jsonsave(str(f)) == 0
    g2 = Grid()
    assert g2.jsonload(str(f)) == 0
    assert g2.cell(0, 0).xstyle == g.cell(0, 0).xstyle
    assert g2.cell(1, 1).type == EMPTY and g2.cell(1, 1).xstyle == XStyle(halign="center")


def test_json_style_from_a_file_is_validated() -> None:
    bad = {"fill": "red", "size": 1e9, "left": {"style": "laser"}, "halign": "middle"}
    assert XStyle.from_json(bad) is None
    assert XStyle.from_json({"fill": "#ff0000", "wrap": True}) == XStyle(fill="FF0000", wrap=True)
    assert XStyle.from_json("nope") is None


def test_far_apart_cells_save_sparsely(tmp_path) -> None:
    g = _excel()
    g.setcell(0, 0, "1")
    g.setcell(NCOL - 1, NROW - 1, "2")
    f = tmp_path / "far.json"
    assert g.jsonsave(str(f)) == 0
    assert f.stat().st_size < 1000
    assert json.loads(f.read_text())["sheets"][0]["cells"] == {"A1": 1, "XFD1048576": 2}
    g2 = Grid()
    assert g2.jsonload(str(f)) == 0
    assert g2.cell(NCOL - 1, NROW - 1).val == 2.0


def test_blanking_a_cell_keeps_its_style() -> None:
    g = _excel()
    g.setcell(0, 0, "5")
    g.cell(0, 0).xstyle = XStyle(fill="FFFF00")
    g.setcell(0, 0, "")
    assert g.cell(0, 0).type == EMPTY and g.cell(0, 0).xstyle == XStyle(fill="FFFF00")


def test_a_spill_keeps_each_cells_own_style() -> None:
    g = _excel()
    g._ensure_cell(0, 2).xstyle = XStyle(fill="00FF00")
    g.setcell(0, 0, "=SEQUENCE(3)")
    assert g.cell(0, 2).type == SPILL and g.cell(0, 2).xstyle == XStyle(fill="00FF00")
    g.setcell(0, 0, "=SEQUENCE(1)")
    assert g.cell(0, 2).type == EMPTY and g.cell(0, 2).xstyle == XStyle(fill="00FF00")


def test_viewport_sends_styles_and_styled_blanks() -> None:
    from gridcalc.web import Api

    g = _excel()
    g.setcell(0, 0, "-5")
    g.cell(0, 0).fmtstr = "0;[Red]-0"
    g._ensure_cell(1, 1).xstyle = XStyle(fill="FFFF00")
    cells = {(x["r"], x["c"]): x for x in Api(g).viewport(0, 0, 3, 3)["cells"]}
    assert cells[(0, 0)]["color"] == "red" and cells[(0, 0)]["text"] == "-5"
    assert cells[(1, 1)] == {"r": 1, "c": 1, "text": "", "align": "l", "style": {"fill": "FFFF00"}}


def test_terminal_view_centres_a_centred_cell() -> None:
    from gridcalc.display import fmtcell

    g = _excel()
    g.setcell(0, 0, "ab")
    g.cell(0, 0).xstyle = XStyle(halign="center")
    assert fmtcell(g.cell(0, 0), 6) == "  ab  "


def test_merges_and_row_heights_round_trip(tmp_path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "title"
    ws.merge_cells("A1:C2")
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[7].height = 12.75
    wb.save(str(tmp_path / "m.xlsx"))
    g = Grid()
    assert g.xlsxload(str(tmp_path / "m.xlsx")) == 0
    assert g._active.merges == [(0, 0, 2, 1)]
    assert g._active.heights == {4: 30.0, 6: 12.75}
    g.mode = Mode.EXCEL
    assert g.xlsxsave(str(tmp_path / "out.xlsx")) == 0
    back = openpyxl.load_workbook(str(tmp_path / "out.xlsx")).active
    assert [str(r) for r in back.merged_cells.ranges] == ["A1:C2"]
    assert (back.row_dimensions[5].height, back.row_dimensions[7].height) == (30, 12.75)
    f = tmp_path / "m.json"
    assert g.jsonsave(str(f)) == 0
    g2 = Grid()
    assert g2.jsonload(str(f)) == 0
    assert (g2._active.merges, g2._active.heights) == ([(0, 0, 2, 1)], {4: 30.0, 6: 12.75})


@pytest.mark.parametrize(
    ("op", "at", "merges", "heights"),
    [
        ("insertrow", 1, [(0, 0, 2, 2)], {5: 30.0}),  # inside: widens
        ("insertrow", 0, [(0, 1, 2, 2)], {5: 30.0}),
        ("deleterow", 1, [(0, 0, 2, 0)], {3: 30.0}),  # inside: narrows
        ("insertcol", 3, [(0, 0, 2, 1)], {4: 30.0}),  # after: unchanged
        ("deletecol", 0, [(0, 0, 1, 1)], {4: 30.0}),
    ],
)
def test_structural_edits_move_merges_and_heights(op, at, merges, heights) -> None:
    g = _excel()
    g.setcell(0, 0, "x")
    g._active.merges = [(0, 0, 2, 1)]
    g._active.heights = {4: 30.0}
    getattr(g, op)(at)
    assert (g._active.merges, g._active.heights) == (merges, heights)


def test_deleting_a_merge_down_to_one_cell_drops_it() -> None:
    g = _excel()
    g._active.merges = [(0, 0, 0, 1)]
    g.deleterow(1)
    assert g._active.merges == []


def test_viewport_sends_merges_and_an_anchor_outside_the_window() -> None:
    from gridcalc.web import Api

    g = _excel()
    g.setcell(0, 0, "Title")
    g._active.merges = [(0, 0, 3, 9)]
    vp = Api(g).viewport(5, 1, 3, 3)  # rows 5-7, cols B-D: inside the merge
    assert vp["merges"] == [[0, 0, 9, 3]]
    assert [(x["r"], x["c"], x["text"]) for x in vp["cells"]] == [(0, 0, "Title")]


def test_undo_keeps_a_styled_blank() -> None:
    from gridcalc.undo import UndoManager

    g = _excel()
    g._ensure_cell(1, 0).xstyle = XStyle(fill="FFFF00")
    u = UndoManager()
    u.save_cell(g, 1, 0)
    g.setcell(1, 0, "5")
    u.undo(g)
    assert g.cell(1, 0).type == EMPTY and g.cell(1, 0).xstyle == XStyle(fill="FFFF00")


def test_a_far_styled_blank_does_not_widen_the_used_extent() -> None:
    g = _excel()
    g.setcell(0, 0, "1")
    g._ensure_cell(5, NROW - 1).xstyle = XStyle(fill="FFFF00")
    assert g._used_extent() == (0, 0)


def test_a_theme_with_a_rels_part_still_resolves(tmp_path) -> None:
    import zipfile

    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.active["A1"].fill = PatternFill("solid", fgColor=Color(theme=4))
    src = tmp_path / "t.xlsx"
    wb.save(str(src))
    with zipfile.ZipFile(src, "a") as z:  # `_rels` sorts before `theme1.xml`
        z.writestr("xl/theme/_rels/theme1.xml.rels", "<Relationships/>")
    g = Grid()
    g.xlsxload(str(src))
    assert g.cell(0, 0).xstyle.fill == "4F81BD"
