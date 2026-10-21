"""Tests for scripts/excel_check.py, which writes the Excel comparison workbook."""

import importlib.util
from pathlib import Path

import pytest

pytest.importorskip("openpyxl")

_SCRIPT = Path(__file__).resolve().parent.parent / "scripts" / "excel_check.py"
_spec = importlib.util.spec_from_file_location("excel_check", _SCRIPT)
excel_check = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(excel_check)


def test_every_case_has_a_row_and_the_inputs_load_as_intended() -> None:
    ws = excel_check.build().active
    assert ws.max_row == len(excel_check.CASES) + 1
    assert ws["I2"].value == "setup" and ws["L2"].value is True


def test_results_keep_their_excel_type() -> None:
    ws = excel_check.build().active
    by_id = {ws[f"I{r}"].value: ws[f"L{r}"] for r in range(2, ws.max_row + 1)}
    assert (by_id["pow-1"].value, by_id["pow-1"].data_type) == (64, "n")
    assert (by_id["text-1"].value, by_id["text-1"].data_type) == ("0.00001", "s")
    assert (by_id["misc-5"].value, by_id["misc-5"].data_type) == ("#N/A", "e")


def test_plain_cases_are_not_array_formulas() -> None:
    ws = excel_check.build().active
    by_id = {ws[f"I{r}"].value: ws[f"K{r}"].value for r in range(2, ws.max_row + 1)}
    assert by_id["misc-7"] == '=OR("a",TRUE)'
    assert isinstance(by_id["misc-3"], excel_check.ArrayFormula)


def test_gridcalc_results_are_stored_to_17_digits(tmp_path) -> None:
    import openpyxl

    path = tmp_path / "check.xlsx"
    excel_check.save(excel_check.build(), path)
    ws = openpyxl.load_workbook(path).active
    by_id = {ws[f"I{r}"].value: ws[f"L{r}"].value for r in range(2, ws.max_row + 1)}
    assert by_id["yf-16"] == 1 / 6  # 16 digits would read back as 0.16666666666666669
    assert by_id["yf-8"] == 366 / 365.5


# Excel's answers from the 2026-10-21 run (docs/dev/excel-check.md), for the
# cases gridcalc has since been changed to match. YEARFRAC values are as the
# run printed them, to 9 decimals.
EXCEL_2026_10_21 = {
    "text-7": "12345678901234500",
    "num-1": 0.0,
    "crit-8": 4.0,
    "crit-9": 3.0,
    "arr-2": "2,2,1",
    "db-4": 0.0,
    "misc-3": True,
    "misc-4": True,
    "misc-6": "1",
    "misc-7": True,
    "misc-8": True,
    "yf-2": 0.166666667,
    "yf-5": 1.0,
    "yf-6": 0.99726776,
    "yf-7": 1.0,
    "yf-10": 0.997260274,
    "yf-11": 0.99726776,
    "yf-12": 0.997260274,
    "yf-14": 1.0,
    "yf-18": 0.086111111,
    "yf-19": 0.083333333,
}


# Exact values from the second 2026-10-21 run, read from the saved workbook:
# the cases new to it, and the ones it showed gridcalc getting wrong.
EXCEL_2026_10_21_XLSX = {
    "text-9": "0.00000000001",
    "text-10": "0.000000000000001",
    "text-15": "0.000000000000000001",
    "text-16": "-0.000000000000000001",
    "text-17": "0.000000000000000015",
    "text-18": "15000000000000000000",
    "text-19": "1.2345E-15",
    "logic-1": True,
    "logic-2": "#VALUE!",
    "logic-3": True,
    "3d-1": 60.0,
    "3d-2": 4.0,
}


@pytest.mark.parametrize("run", [EXCEL_2026_10_21, EXCEL_2026_10_21_XLSX])
def test_gridcalc_agrees_with_excel_on_the_fixed_cases(run: dict[str, object]) -> None:
    ids = [cid for cid, _, _ in excel_check.CASES]
    ours = dict(zip(ids, excel_check.evaluate(), strict=True))
    for cid, excel in run.items():
        if isinstance(excel, float):
            assert ours[cid] == pytest.approx(excel, abs=6e-10), cid
        else:
            assert ours[cid] == excel, cid


def test_3d_cases_read_the_span_sheets_as_plain_formulas() -> None:
    wb = excel_check.build()
    assert wb.sheetnames == ["check", "S3a", "S3b", "S3c"]
    ws = wb["check"]
    rows = {ws[f"I{r}"].value: r for r in range(2, ws.max_row + 1)}
    assert ws[f"K{rows['3d-1']}"].value == "=SUM(S3a:S3c!B2)"  # not an ArrayFormula
    assert (ws[f"L{rows['3d-1']}"].value, ws[f"L{rows['3d-2']}"].value) == (60.0, 4.0)
