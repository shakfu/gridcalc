"""Defined names holding a constant, a formula or a LAMBDA (Excel's Name Manager)."""

from pathlib import Path

import pytest

from gridcalc import commands as shared
from gridcalc.engine import Grid, Mode, NamedRange
from gridcalc.formula.errors import ExcelError
from gridcalc.undo import UndoManager


def _grid(**names: str) -> Grid:
    g = Grid()
    g.mode = Mode.EXCEL
    g._apply_mode_libs()
    g.names = [NamedRange(n, formula=f) for n, f in names.items()]
    return g


def test_a_constant_name() -> None:
    g = _grid(Rate="=0.2")
    g.setcell(0, 0, "100")
    g.setcell(1, 0, "=A1*Rate")
    assert g.cells[1][0].val == 20.0


def test_a_formula_name_recalculates_when_its_inputs_change() -> None:
    g = _grid(Total="=SUM(Sheet1!$A$1:$A$3)")
    g.setcells_bulk([(0, 0, "1"), (0, 1, "2"), (1, 5, "=Total*2")])
    assert g.cells[1][5].val == 6.0
    g.setcell(0, 2, "10")
    assert g.cells[1][5].val == 26.0


def test_a_volatile_name_makes_its_cell_volatile() -> None:
    g = _grid(Next="=OFFSET(Sheet1!$A$1,1,0)")
    g.setcells_bulk([(0, 0, "1"), (0, 1, "2"), (1, 0, "=Next")])
    assert g.cells[1][0].val == 2.0
    g.setcell(0, 1, "7")
    assert g.cells[1][0].val == 7.0


def test_a_named_lambda_is_callable_and_recurses() -> None:
    g = _grid(
        Double="=LAMBDA(x, x*2)",
        Fact="=LAMBDA(n, IF(n<2, 1, n*Fact(n-1)))",
    )
    g.setcell(0, 0, "=Double(4)+Fact(5)")
    assert g.cells[0][0].val == 128.0


def test_a_named_lambda_reads_cells_in_its_body() -> None:
    g = _grid(Scaled="=LAMBDA(x, x*Sheet1!$C$1)")
    g.setcells_bulk([(2, 0, "3"), (0, 0, "=Scaled(2)")])
    assert g.cells[0][0].val == 6.0
    g.setcell(2, 0, "5")
    assert g.cells[0][0].val == 10.0


def test_a_name_defined_in_terms_of_itself_is_an_error() -> None:
    g = _grid(Loop="=Loop+1", A="=B", B="=A")
    g.setcell(0, 0, "=Loop")
    g.setcell(0, 1, "=A")
    assert g.cells[0][0].err is ExcelError.CIRC
    assert g.cells[0][1].err is ExcelError.CIRC


def test_row_insert_and_sheet_rename_rewrite_a_name_formula() -> None:
    g = _grid(Total="=SUM(Sheet1!$A$2:$A$3)")
    g.insertrow(0)
    assert g.names[0].formula == "=SUM(Sheet1!$A$3:$A$4)"
    g.add_sheet("Other")
    g.rename_sheet("Sheet1", "Data")
    assert g.names[0].formula == "=SUM(Data!$A$3:$A$4)"
    g.remove_sheet("Data")
    assert g.names[0].formula == "=SUM(#REF!)"


def test_json_roundtrip(tmp_path: Path) -> None:
    g = _grid(Rate="=0.2")
    g.setcell(0, 0, "=Rate*10")
    f = tmp_path / "book.json"
    assert g.jsonsave(str(f)) == 0
    g2 = Grid()
    assert g2.jsonload(str(f)) == 0
    assert [(n.name, n.formula) for n in g2.names] == [("Rate", "=0.2")]
    assert g2.cells[0][0].val == 2.0


def test_xlsx_roundtrip(tmp_path: Path) -> None:
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.workbook.defined_name import DefinedName

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = 4
    ws["B1"] = "=Rate*A1"
    ws["B2"] = "=Sq(3)"
    wb.defined_names["Rate"] = DefinedName("Rate", attr_text="0.5")
    wb.defined_names["Sq"] = DefinedName("Sq", attr_text="_xlfn.LAMBDA(_xlpm.x,_xlpm.x*_xlpm.x)")
    src = tmp_path / "in.xlsx"
    wb.save(str(src))
    g = Grid()
    assert g.xlsxload(str(src)) == 0
    assert {n.name: n.formula for n in g.names} == {"Rate": "=0.5", "Sq": "=LAMBDA(x,x*x)"}
    assert (g.cells[1][0].val, g.cells[1][1].val) == (2.0, 9.0)
    out = tmp_path / "out.xlsx"
    assert g.xlsxsave(str(out)) == 0
    back = openpyxl.load_workbook(str(out))
    assert back.defined_names["Sq"].attr_text == "_xlfn.LAMBDA(_xlpm.x,_xlpm.x*_xlpm.x)"
    assert back.defined_names["Rate"].attr_text == "0.5"


def test_name_command_defines_and_lists_a_formula() -> None:
    g = _grid()
    g.setcell(0, 0, "10")
    g.setcell(1, 0, "=A1*Rate")
    res = shared.run("name", g, UndoManager(), ["Rate", "=1", "+", "1"], None)
    assert res.ok and res.message == "Rate = =1 + 1"
    assert g.cells[1][0].val == 20.0
    listed = shared.run("names", g, UndoManager(), [], None)
    assert listed.message == "Rate==1 + 1"
    assert not shared.run("name", g, UndoManager(), ["Bad", "=1+"], None).ok


def test_a_saved_name_that_reads_as_a_cell_warns_on_load(tmp_path: Path) -> None:
    f = tmp_path / "old.json"
    f.write_text('{"version": 2, "mode": "EXCEL", "names": {"TAX1": "B2:B5"}, "cells": [[1]]}')
    g = Grid()
    assert g.jsonload(str(f)) == 0
    assert g.load_warnings == [
        "name TAX1 reads as a cell reference, since columns now reach XFD; rename it"
    ]
