import math
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")

from gridcalc.engine import Grid, Mode, NamedRange  # noqa: E402

EXAMPLES = Path(__file__).resolve().parent.parent / "examples"


def test_xlsxsave_empty_grid_returns_minus_one(tmp_path):
    g = Grid()
    assert g.xlsxsave(str(tmp_path / "empty.xlsx")) == -1


def test_xlsxload_multisheet_preserves_sheets(tmp_path):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Inputs"
    ws1.cell(row=1, column=1, value=10)
    ws1.cell(row=2, column=1, value=20)
    ws2 = wb.create_sheet("Outputs")
    ws2.cell(row=1, column=1, value=42)
    ws2.cell(row=2, column=1, value="hello")
    f = tmp_path / "multi.xlsx"
    wb.save(str(f))

    g = Grid()
    assert g.xlsxload(str(f)) == 0
    assert g.sheet_names() == ["Inputs", "Outputs"]
    # First xlsx sheet becomes active.
    assert g.active == 0
    assert g.cells[0][0].val == 10.0
    assert g.cells[0][1].val == 20.0
    g.set_active("Outputs")
    assert g.cells[0][0].val == 42.0
    assert g.cells[0][1].text == "hello"


def test_xlsxsave_multisheet_writes_all_sheets(tmp_path):
    g = Grid()
    g.mode = Mode.EXCEL
    g._apply_mode_libs()
    g.rename_sheet("Sheet1", "Inputs")
    g.setcell(0, 0, "10")
    g.add_sheet("Outputs")
    g.set_active("Outputs")
    g.setcell(0, 0, "20")
    f = tmp_path / "saved.xlsx"
    assert g.xlsxsave(str(f)) == 0

    wb = openpyxl.load_workbook(str(f))
    assert set(wb.sheetnames) == {"Inputs", "Outputs"}
    assert wb["Inputs"].cell(row=1, column=1).value == 10.0
    assert wb["Outputs"].cell(row=1, column=1).value == 20.0


def test_xlsx_multisheet_roundtrip_preserves_cross_sheet_formula_value(tmp_path):
    g = Grid()
    g.mode = Mode.EXCEL
    g._apply_mode_libs()
    g.rename_sheet("Sheet1", "Source")
    g.setcell(0, 0, "10")
    g.add_sheet("Derived")
    g.set_active("Derived")
    g.setcell(0, 0, "=Source!A1*5")
    assert g.cells[0][0].val == 50.0
    f = tmp_path / "rt.xlsx"
    assert g.xlsxsave(str(f)) == 0

    g2 = Grid()
    assert g2.xlsxload(str(f)) == 0
    assert set(g2.sheet_names()) == {"Source", "Derived"}
    g2.set_active("Source")
    assert g2.cells[0][0].val == 10.0
    g2.set_active("Derived")
    # In EXCEL mode the formula text round-trips, and gridcalc
    # re-evaluates it from the live Source!A1.
    assert g2.cells[0][0].text == "=Source!A1*5"
    assert g2.cells[0][0].val == 50.0
    # And the file itself carries the formula (data_only=False) plus
    # the cached numeric value (data_only=True).
    wb_f = openpyxl.load_workbook(str(f), data_only=False)
    assert wb_f["Derived"].cell(row=1, column=1).value == "=Source!A1*5"
    wb_v = openpyxl.load_workbook(str(f), data_only=True)
    assert wb_v["Derived"].cell(row=1, column=1).value == 50.0


def test_xlsxsave_basic(tmp_path):
    g = Grid()
    g.setcell(0, 0, "Header")
    g.setcell(1, 0, "10")
    g.setcell(2, 0, "=B1*2")
    f = tmp_path / "out.xlsx"
    assert g.xlsxsave(str(f)) == 0
    wb = openpyxl.load_workbook(str(f), data_only=False)
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Header"
    assert ws.cell(row=1, column=2).value == 10.0
    # legacy mode evaluates =B1*2 = 20; we save the value
    assert ws.cell(row=1, column=3).value == 20.0


def test_xlsxload_values(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="City")
    ws.cell(row=1, column=2, value="Pop")
    ws.cell(row=2, column=1, value="NYC")
    ws.cell(row=2, column=2, value=8000000)
    f = tmp_path / "data.xlsx"
    wb.save(str(f))

    g = Grid()
    assert g.xlsxload(str(f)) == 0
    assert g.mode == Mode.EXCEL
    assert "xlsx" in g.libs
    assert g.cells[0][0].text == "City"
    assert g.cells[1][1].val == 8000000.0


def test_xlsxload_formulas(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value=10)
    ws.cell(row=1, column=2, value=20)
    ws.cell(row=1, column=3, value="=A1+B1")
    f = tmp_path / "f.xlsx"
    wb.save(str(f))

    g = Grid()
    assert g.xlsxload(str(f)) == 0
    # formula should be re-evaluated by gridcalc EXCEL evaluator
    assert g.cells[2][0].val == 30.0


def test_xlsxload_then_save_roundtrip(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value=5)
    ws.cell(row=1, column=2, value="=A1*3")
    src = tmp_path / "src.xlsx"
    wb.save(str(src))

    g = Grid()
    assert g.xlsxload(str(src)) == 0
    dst = tmp_path / "dst.xlsx"
    assert g.xlsxsave(str(dst)) == 0
    # In EXCEL mode, xlsxsave preserves formula text and writes the
    # cached numeric value. data_only=False returns the formula string;
    # data_only=True returns the cached value.
    wb2 = openpyxl.load_workbook(str(dst), data_only=False)
    ws2 = wb2.active
    assert ws2.cell(row=1, column=1).value == 5.0
    assert ws2.cell(row=1, column=2).value == "=A1*3"
    wb2v = openpyxl.load_workbook(str(dst), data_only=True)
    assert wb2v.active.cell(row=1, column=2).value == 15.0


def test_xlsxload_bool(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value=True)
    ws.cell(row=1, column=2, value=False)
    f = tmp_path / "b.xlsx"
    wb.save(str(f))

    g = Grid()
    assert g.xlsxload(str(f)) == 0
    assert g.cells[0][0].val is True and g.cells[0][0].sval == "TRUE"
    assert g.cells[1][0].val is False and g.cells[1][0].sval == "FALSE"


def test_xlsxload_unknown_function_yields_nan(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="=NONESUCH()")
    f = tmp_path / "u.xlsx"
    wb.save(str(f))

    g = Grid()
    assert g.xlsxload(str(f)) == 0
    assert math.isnan(g.cells[0][0].val)


def test_xlsxload_xlsx_lib_functions_work(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value=1)
    ws.cell(row=2, column=1, value=2)
    ws.cell(row=3, column=1, value=3)
    ws.cell(row=4, column=1, value="=AVERAGE(A1:A3)")
    f = tmp_path / "avg.xlsx"
    wb.save(str(f))

    g = Grid()
    assert g.xlsxload(str(f)) == 0
    assert g.cells[0][3].val == 2.0


def test_xlsxload_missing_file_returns_minus_one(tmp_path):
    g = Grid()
    assert g.xlsxload(str(tmp_path / "nope.xlsx")) == -1


def test_example_multisheet_xlsx_roundtrip_preserves_formulas(tmp_path):
    """Import examples/example_multisheet.xlsx and re-export it; the
    written file must be a well-formed xlsx that preserves both the
    formula text and a cached numeric value (where computable)."""
    src = EXAMPLES / "example_multisheet.xlsx"
    assert src.is_file(), f"missing example fixture: {src}"

    g = Grid()
    assert g.xlsxload(str(src)) == 0
    assert g.mode == Mode.EXCEL
    assert g.sheet_names() == ["Jan", "Feb", "Summary"]

    # Sanity-check imported content.
    g.set_active("Jan")
    assert g.cells[1][1].val == 100.0
    g.set_active("Feb")
    assert g.cells[1][2].val == 130.0
    g.set_active("Summary")
    assert g.cells[0][1].text == "=SUM(Jan!B2:B3)+SUM(Feb!B2:B3)"
    expected_total = 100.0 + 150.0 + 120.0 + 130.0
    assert g.cells[0][1].val == expected_total

    dst = tmp_path / "rt_example.xlsx"
    assert g.xlsxsave(str(dst)) == 0

    # Well-formed: openpyxl can open it and find every sheet/cell.
    wb_f = openpyxl.load_workbook(str(dst), data_only=False)
    assert wb_f.sheetnames == ["Jan", "Feb", "Summary"]
    assert wb_f["Jan"].cell(row=1, column=1).value == "Item"
    assert wb_f["Jan"].cell(row=2, column=2).value == 100.0
    assert wb_f["Feb"].cell(row=3, column=2).value == 130.0
    # Formula text round-trips as a formula, not as a stringified value.
    summary_a2 = wb_f["Summary"].cell(row=2, column=1)
    assert summary_a2.data_type == "f"
    assert summary_a2.value == "=SUM(Jan!B2:B3)+SUM(Feb!B2:B3)"

    # Cached value is also written so data_only readers see a number.
    wb_v = openpyxl.load_workbook(str(dst), data_only=True)
    assert wb_v["Summary"].cell(row=2, column=1).value == expected_total


def test_xlsxsave_keeps_default_sheet1_alongside_another_sheet(tmp_path):
    # The default sheet name is the one nobody renames. Because OpenXLSX also
    # calls its auto-created sheet "Sheet1", the export used to treat the first
    # payload sheet as "already present", leave the default unclaimed, and then
    # rename it to the *second* sheet -- merging both sheets into one.
    g = Grid()
    g.mode = Mode.EXCEL
    g._apply_mode_libs()
    assert g.sheet_names() == ["Sheet1"]
    g.setcell(0, 0, "111")
    g.add_sheet("Data")
    g.set_active("Data")
    g.setcell(0, 0, "222")
    f = tmp_path / "default_first.xlsx"
    assert g.xlsxsave(str(f)) == 0

    wb = openpyxl.load_workbook(str(f))
    assert wb.sheetnames == ["Sheet1", "Data"]
    assert wb["Sheet1"].cell(row=1, column=1).value == 111.0
    assert wb["Data"].cell(row=1, column=1).value == 222.0


def test_xlsxsave_writes_empty_sheets(tmp_path):
    # A sheet with no cells contributes nothing to the cell payload, so the
    # writer needs the workbook's sheet list to know it exists at all.
    g = Grid()
    g.mode = Mode.EXCEL
    g._apply_mode_libs()
    g.rename_sheet("Sheet1", "Inputs")
    g.setcell(0, 0, "1")
    g.add_sheet("Blank")
    g.add_sheet("Outputs")
    g.set_active("Outputs")
    g.setcell(0, 0, "2")
    f = tmp_path / "with_empty.xlsx"
    assert g.xlsxsave(str(f)) == 0

    wb = openpyxl.load_workbook(str(f))
    assert wb.sheetnames == ["Inputs", "Blank", "Outputs"]
    assert wb["Blank"].max_row == 1
    assert wb["Blank"].cell(row=1, column=1).value is None


def test_xlsxsave_preserves_sheet_order_when_first_sheet_is_empty(tmp_path):
    # Sheet order must follow the workbook model, not whichever sheet holds
    # the first non-empty cell.
    g = Grid()
    g.mode = Mode.EXCEL
    g._apply_mode_libs()
    g.rename_sheet("Sheet1", "Cover")
    g.add_sheet("Body")
    g.set_active("Body")
    g.setcell(0, 0, "7")
    f = tmp_path / "order.xlsx"
    assert g.xlsxsave(str(f)) == 0

    wb = openpyxl.load_workbook(str(f))
    assert wb.sheetnames == ["Cover", "Body"]
    assert wb["Body"].cell(row=1, column=1).value == 7.0


# -- defined names --


def _names(g):
    return sorted((n.name, n.sheet, n.c1, n.r1, n.c2, n.r2) for n in g.names)


def test_example_excel_xlsx_matches_json():
    """The shipped xlsx is example_excel.json saved as xlsx; both evaluate alike."""
    from gridcalc.display import cell_text, fmtcell
    from gridcalc.loader import load_workbook

    a = load_workbook(EXAMPLES / "example_excel.json")
    b = load_workbook(EXAMPLES / "example_excel.xlsx")
    assert _names(a) == _names(b)
    assert a.cw == b.cw == 12
    assert any(cl.fmtstr for cl in a._cells.values())  # the example carries number formats
    for pos, cl in a._cells.items():
        assert cell_text(cl) == cell_text(b._cells[pos]), pos
        assert fmtcell(cl, a.cw, a.fmt) == fmtcell(b._cells[pos], b.cw, b.fmt), pos
        assert bool(cl.bold) == bool(b._cells[pos].bold), pos


def test_xlsxsave_writes_names_other_readers_see(tmp_path):
    g = Grid()
    g.mode = Mode.EXCEL
    g._apply_mode_libs()
    g.setcell(0, 0, "1")
    g.add_sheet("My 'Data")
    g.set_active("My 'Data")
    g.setcell(1, 1, "2")
    g.set_active(0)
    g.setcell(0, 1, "=SUM(both)")
    g.names = [
        NamedRange(name="both", c1=0, r1=0, c2=0, r2=0),
        NamedRange(name="data", c1=1, r1=1, c2=1, r2=2, sheet="My 'Data"),
    ]
    f = tmp_path / "names.xlsx"
    assert g.xlsxsave(str(f)) == 0

    wb = openpyxl.load_workbook(str(f))
    assert wb.defined_names["data"].attr_text == "'My ''Data'!$B$2:$B$3"
    # A sheet-less name is local on each sheet, pointing at that sheet.
    assert wb["Sheet1"].defined_names["both"].attr_text == "'Sheet1'!$A$1"
    assert wb["My 'Data"].defined_names["both"].attr_text == "'My ''Data'!$A$1"


def test_xlsx_names_round_trip(tmp_path):
    g = Grid()
    g.mode = Mode.EXCEL
    g._apply_mode_libs()
    g.setcell(0, 0, "5")
    g.add_sheet("S2")
    g.set_active("S2")
    g.setcell(0, 0, "7")
    g.setcell(1, 0, "=x*2")
    g.set_active(0)
    g.setcell(1, 0, "=x*2")
    g.names = [
        NamedRange(name="x", c1=0, r1=0, c2=0, r2=0),
        NamedRange(name="y", c1=0, r1=0, c2=0, r2=3, sheet="S2"),
    ]
    g.recalc()
    f = tmp_path / "rt.xlsx"
    assert g.xlsxsave(str(f)) == 0

    g2 = Grid()
    assert g2.xlsxload(str(f)) == 0
    assert _names(g2) == _names(g)
    g2.recalc()
    # `x` still resolves on each formula's own sheet.
    assert g2.sheets[0]._cells[(1, 0)].val == 10
    assert g2.sheets[1]._cells[(1, 0)].val == 14


def test_partial_local_names_stay_sheet_qualified(tmp_path):
    """A name local to only some sheets is not a sheet-less name."""
    wb = openpyxl.Workbook()
    wb.active.title = "A"
    wb.create_sheet("B")
    from openpyxl.workbook.defined_name import DefinedName

    wb["A"].defined_names["loc"] = DefinedName("loc", localSheetId=0, attr_text="A!$A$1")
    f = tmp_path / "local.xlsx"
    wb.save(str(f))

    g = Grid()
    assert g.xlsxload(str(f)) == 0
    assert _names(g) == [("loc", "A", 0, 0, 0, 0)]


def test_save_losses_keeps_names_for_xlsx():
    from gridcalc.loader import save_losses

    g = Grid()
    g.setcell(0, 0, "1")
    g.names = [NamedRange(name="n", c1=0, r1=0, c2=0, r2=0)]
    assert "names" not in save_losses(g, "out.xlsx")
    assert "names" in save_losses(g, "out.csv")


# -- column widths --


def test_xlsxsave_writes_column_widths_other_readers_see(tmp_path):
    g = Grid()
    g.mode = Mode.EXCEL
    g._apply_mode_libs()
    g.cw = 12
    g.setcell(0, 0, "1")
    g._active.widths = {1: 110}
    f = tmp_path / "w.xlsx"
    assert g.xlsxsave(str(f)) == 0

    ws = openpyxl.load_workbook(str(f))["Sheet1"]
    assert ws.sheet_format.defaultColWidth == 12
    assert ws.column_dimensions["B"].width == pytest.approx((110 - 5) / 7, abs=0.01)


def test_xlsx_column_widths_round_trip(tmp_path):
    g = Grid()
    g.mode = Mode.EXCEL
    g._apply_mode_libs()
    g.cw = 15
    g.setcell(0, 0, "1")
    g._active.widths = {0: 80, 3: 240}
    g.add_sheet("S2")
    g.sheets[1].widths = {2: 60}
    f = tmp_path / "rt.xlsx"
    assert g.xlsxsave(str(f)) == 0

    g2 = Grid()
    assert g2.xlsxload(str(f)) == 0
    assert g2.cw == 15
    assert g2.sheets[0].widths == {0: 80, 3: 240}
    assert g2.sheets[1].widths == {2: 60}


def test_xlsxload_reads_widths_from_other_writers(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws.sheet_format.defaultColWidth = 20
    ws.column_dimensions["C"].width = 25
    f = tmp_path / "o.xlsx"
    wb.save(str(f))

    g = Grid()
    assert g.xlsxload(str(f)) == 0
    assert g.cw == 20
    assert g._active.widths == {2: 25 * 7 + 5}


# -- text style --


def _styled_grid():
    g = Grid()
    g.mode = Mode.EXCEL
    g._apply_mode_libs()
    g.setcell(0, 0, '"Head')
    g.setcell(1, 0, "2")
    g.setcell(2, 0, "=B1*2")
    g.setcell(3, 0, '"plain')
    g.cell(0, 0).bold = 1
    g.cell(1, 0).italic = 1
    g.cell(2, 0).underline = 1
    g.cell(2, 0).bold = 1
    g.recalc()
    return g


def _style(cl):
    return (cl.bold, cl.italic, cl.underline)


def test_xlsxsave_writes_text_style_other_readers_see(tmp_path):
    f = tmp_path / "s.xlsx"
    assert _styled_grid().xlsxsave(str(f)) == 0
    ws = openpyxl.load_workbook(str(f))["Sheet1"]
    assert ws["A1"].font.b and not ws["A1"].font.i
    assert ws["B1"].font.i and not ws["B1"].font.b
    assert ws["C1"].font.b and ws["C1"].font.u == "single"
    assert not ws["D1"].font.b and ws["D1"].font.u is None


def test_xlsx_text_style_round_trip(tmp_path):
    g = _styled_grid()
    f = tmp_path / "s.xlsx"
    assert g.xlsxsave(str(f)) == 0
    g2 = Grid()
    assert g2.xlsxload(str(f)) == 0
    for c in range(4):
        assert _style(g2.cell(c, 0)) == _style(g.cell(c, 0)), c


def test_xlsx_style_shares_one_format_per_style(tmp_path):
    """17 bold cells need one bold font and one cell format, not 17."""
    g = Grid()
    for r in range(17):
        g.setcell(0, r, '"h')
        g.cell(0, r).bold = 1
    f = tmp_path / "many.xlsx"
    assert g.xlsxsave(str(f)) == 0
    wb = openpyxl.load_workbook(str(f))
    xfs = {wb["Sheet1"].cell(row=r + 1, column=1).style_id for r in range(17)}
    assert len(xfs) == 1


def test_xlsxload_reads_style_from_other_writers(tmp_path):
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "x"
    ws["A1"].font = Font(bold=True, italic=True, underline="double")
    f = tmp_path / "o.xlsx"
    wb.save(str(f))
    g = Grid()
    assert g.xlsxload(str(f)) == 0
    assert _style(g.cell(0, 0)) == (1, 1, 1)


# -- number formats and alignment --


@pytest.mark.parametrize(
    "spec, code",
    [
        (",.2f", "#,##0.00"),
        (".2f", "0.00"),
        (",", "#,##0"),
        (".0f", "0"),
        (".1%", "0.0%"),
        (".2e", "0.00E+00"),
        ("f", "0.000000"),
    ],
)
def test_spec_and_xlsx_code_convert_both_ways(spec, code):
    from gridcalc.display import fmt_float
    from gridcalc.engine import _spec_for_xlsx_code, _xlsx_code_for_spec

    assert _xlsx_code_for_spec(spec) == code
    back = _spec_for_xlsx_code(code, 0)
    assert fmt_float(1234.5678, back) == fmt_float(1234.5678, spec)


def test_unmapped_formats_convert_to_nothing():
    from gridcalc.engine import _spec_for_xlsx_code, _xlsx_code_for_spec

    assert _xlsx_code_for_spec("x>10") == ""
    assert _spec_for_xlsx_code('"$"#,##0.00', 0) == ""
    assert _spec_for_xlsx_code("General", 0) == ""
    assert _spec_for_xlsx_code("", 10) == ".2%"  # built-in id, no code


def _formatted_grid():
    g = Grid()
    g.mode = Mode.EXCEL
    g._apply_mode_libs()
    cases = [
        ("1234.5", "$", ""),
        ("0.157", "%", ""),
        ("42", "I", ""),
        ("7", "L", ""),
        ("7", "R", ""),
        ("1234.5", "", ",.2f"),
        ("0.157", "", ".1%"),
        ("12345", "", ".2e"),
        ("46000", "", "yyyy-mm-dd"),
        ('=CONCAT("a","b")', "L", ""),
        ('=CONCAT("a","b")', "", ""),
    ]
    for r, (text, fmt, fmtstr) in enumerate(cases):
        g.setcell(0, r, text)
        g.cell(0, r).fmt = fmt
        g.cell(0, r).fmtstr = fmtstr
    g.recalc()
    return g, len(cases)


def test_xlsx_number_formats_round_trip_as_displayed(tmp_path):
    from gridcalc.display import fmtcell

    g, n = _formatted_grid()
    f = tmp_path / "fmt.xlsx"
    assert g.xlsxsave(str(f)) == 0
    g2 = Grid()
    assert g2.xlsxload(str(f)) == 0
    g2.recalc()
    for r in range(n):
        assert fmtcell(g2.cell(0, r), 12, g2.fmt) == fmtcell(g.cell(0, r), 12, g.fmt), r


def test_xlsxsave_writes_formats_other_readers_see(tmp_path):
    g, _ = _formatted_grid()
    g.fmt = "$"  # the workbook default reaches cells with no letter of their own
    g.setcell(1, 0, "3")
    f = tmp_path / "fmt.xlsx"
    assert g.xlsxsave(str(f)) == 0
    ws = openpyxl.load_workbook(str(f))["Sheet1"]
    assert ws["A1"].number_format == "0.00"
    assert ws["A2"].number_format == "0.00%"
    assert ws["A6"].number_format == "#,##0.00"
    assert ws["A4"].alignment.horizontal == "left"
    assert ws["A5"].alignment.horizontal == "right"
    assert ws["A11"].alignment.horizontal == "right"
    assert ws["B1"].number_format == "0.00"


def test_bar_format_has_no_xlsx_equivalent(tmp_path):
    g = Grid()
    g.setcell(0, 0, "5")
    g.cell(0, 0).fmt = "*"
    f = tmp_path / "bar.xlsx"
    assert g.xlsxsave(str(f)) == 0
    assert openpyxl.load_workbook(str(f))["Sheet1"]["A1"].number_format == "General"


def test_xlsxload_reads_formats_from_other_writers(tmp_path):
    from openpyxl.styles import Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 0.25
    ws["A1"].number_format = "0.00%"
    ws["A2"] = 1234.5
    ws["A2"].number_format = "#,##0.00"
    ws["A3"] = 9
    ws["A3"].alignment = Alignment(horizontal="left")
    f = tmp_path / "o.xlsx"
    wb.save(str(f))
    g = Grid()
    assert g.xlsxload(str(f)) == 0
    assert g.cell(0, 0).fmtstr == ".2%"
    assert g.cell(0, 1).fmtstr == ",.2f"
    assert g.cell(0, 2).fmt == "L"


# -- saving over an xlsx: what gridcalc cannot write back --


def _rich_xlsx(path):
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Head"
    ws["A1"].font = Font(size=16, color="FF0000")
    ws["A1"].comment = Comment("note", "me")
    ws["B1"] = 1
    ws["B1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["B1"].border = Border(top=Side("thin"))
    ws["C1"] = 2
    ws["C1"].alignment = Alignment(horizontal="center")
    ws["D1"] = 3
    ws["D1"].number_format = '"$"#,##0.00'
    ws.merge_cells("A3:B3")
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    wb.save(str(path))


def test_xlsx_extras_names_what_a_save_would_drop(tmp_path):
    from gridcalc.loader import xlsx_extras

    f = tmp_path / "rich.xlsx"
    _rich_xlsx(f)
    assert set(xlsx_extras(f)) == {
        "comments",
        "merged cells",
        "frozen panes",
        "row heights",
        "font sizes, colours or faces",
        "cell fills",
        "borders",
        "centred, vertical or wrapped alignment",
        "number formats such as currency",
    }


def test_gridcalc_output_has_no_extras(tmp_path):
    """Our own files must not trigger the question, or every `:w` would ask."""
    from gridcalc.loader import xlsx_extras

    blank = tmp_path / "blank.xlsx"
    openpyxl.Workbook().save(str(blank))
    assert xlsx_extras(blank) == []
    for name, g in (("styled", _styled_grid()), ("formatted", _formatted_grid()[0])):
        g.names = [NamedRange(name="n", c1=0, r1=0, c2=0, r2=0)]
        g._active.widths = {1: 100}
        f = tmp_path / f"{name}.xlsx"
        assert g.xlsxsave(str(f)) == 0
        assert xlsx_extras(f) == [], name
    assert xlsx_extras(EXAMPLES / "example_excel.xlsx") == []


def test_xlsx_extras_of_an_unreadable_file_is_empty(tmp_path):
    from gridcalc.loader import xlsx_extras

    bad = tmp_path / "bad.xlsx"
    bad.write_text("not a zip")
    assert xlsx_extras(bad) == []
    assert xlsx_extras(tmp_path / "missing.xlsx") == []


def test_save_losses_includes_the_target_files_extras(tmp_path):
    from gridcalc.loader import save_losses

    f = tmp_path / "rich.xlsx"
    _rich_xlsx(f)
    g = Grid()
    g.mode = Mode.EXCEL
    g.setcell(0, 0, "1")
    assert "merged cells" in save_losses(g, f)
    assert save_losses(g, tmp_path / "new.xlsx") == []
